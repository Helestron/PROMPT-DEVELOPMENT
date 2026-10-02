"""Testes dos scripts do skill gabinete-lessa-criminal-tjal (rodar com: python -m pytest tests)."""
import json
import subprocess
import sys
from pathlib import Path

import pytest

SKILL = Path(__file__).resolve().parent.parent / "skills" / "gabinete-lessa-criminal-tjal"
sys.path.insert(0, str(SKILL / "scripts"))

import cnj  # noqa: E402
import conferir_citacoes  # noqa: E402
import dosimetria  # noqa: E402
import gerar_versoes  # noqa: E402
import montar_minuta  # noqa: E402
import prescricao  # noqa: E402
import regimento  # noqa: E402
import verificar_minuta  # noqa: E402

VOTO_OK = """# RELATÓRIO
Trata-se de apelação criminal interposta por **J. S.** contra a sentença de fls. 210/218, que o condenou pela prática do crime previsto no art. 33, caput, da Lei n.º 11.343/2006.
Nas razões de fls. 230/241, a defesa pede a aplicação da causa de diminuição do art. 33, § 4º, da Lei n.º 11.343/2006, e o Ministério Público, nas contrarrazões de fls. 245/250, pugna pelo desprovimento.
É o relatório.
# VOTO
Presentes os pressupostos de admissibilidade, conheço do recurso e passo ao exame do **pedido de aplicação do tráfico privilegiado** {{v: conferir certidão de antecedentes de fls. 80}}, à luz do HC 598.051/SP.
Desse modo, a primariedade comprovada às fls. 80 e a ausência de prova de dedicação a atividades criminosas autorizam o reconhecimento da minorante, na fração máxima, como pleiteado pela defesa.
Ante o exposto, voto no sentido de conhecer do recurso e dar-lhe provimento, para reconhecer a causa de diminuição e redimensionar a pena nos termos da fundamentação.
É como voto.
"""


# ---------- cnj ----------
def test_cnj_digito_e_validacao():
    d = cnj.digito("0700123", "2024", "8", "02", "0001")
    assert cnj.analisar(f"0700123-{d}.2024.8.02.0001")["valido"]
    errado = cnj.analisar("0700123-00.2024.8.02.0001")
    assert not errado["valido"] and "dígito" in errado["motivo"]


def test_cnj_originario_e_completar():
    n = cnj.completar("0800001-00.2026", "0000")
    assert n.endswith(".8.02.0000") and cnj.analisar(n)["originario_2grau"]
    assert cnj.extrair("autos 0700123-83.2024.8.02.0001 e 0700123-83.2024.8.02.0001") == ["0700123-83.2024.8.02.0001"]


# ---------- dosimetria ----------
def caso_trafico(**afirmado):
    return {"pena_min": "5a", "pena_max": "15a", "multa_min": 500, "multa_max": 1500,
            "fase1": {"criterio": "1/8_intervalo", "desfavoraveis": 1},
            "fase2": {"agravantes": 0, "atenuantes": 1, "fracao": "1/6"},
            "fase3": [{"tipo": "diminuicao", "fracao": "2/3", "fundamento": "art. 33, § 4º"}],
            "afirmado": afirmado}


def test_dosimetria_confere():
    r = dosimetria.calcular(caso_trafico(pena_final="1a8m25d", multa_final=173))
    assert r["conclusao"] == "CONFERE"
    assert r["pena_final"] == "1 ano, 8 meses e 25 dias"


def test_dosimetria_diverge():
    r = dosimetria.calcular(caso_trafico(pena_final="1a8m"))
    assert r["conclusao"] == "DIVERGE"


def test_dosimetria_sumula231_limita_no_minimo():
    e = caso_trafico()
    e["fase1"]["desfavoraveis"] = 0
    e["fase3"] = []
    r = dosimetria.calcular(e)
    assert r["pena_final"] == "5 anos"


# ---------- prescrição ----------
def test_prescricao_art10_inclui_dia_do_comeco():
    base = {"pena": "1a4m", "modalidade": "concreto", "transito_acusacao": True,
            "marcos": [{"evento": "recebimento da denúncia", "data": "2018-08-01"}]}
    no_limite = prescricao.calcular({**base, "data_referencia": "2022-07-31"})
    passou = prescricao.calcular({**base, "data_referencia": "2022-08-01"})
    assert no_limite["conclusao"] == "NAO_CONSUMADA"
    assert passou["conclusao"] == "PRESCRICAO_CONSUMADA"


def test_prescricao_reducao_art115_e_suspensao():
    e = {"pena": "1a4m", "modalidade": "concreto", "transito_acusacao": True, "menor_21_no_fato": True,
         "marcos": [{"evento": "recebimento da denúncia", "data": "2020-01-01"}],
         "suspensoes": [{"inicio": "2020-06-01", "fim": "2020-12-31"}],
         "data_referencia": "2022-03-01"}
    r = prescricao.calcular(e)
    assert r["prazo_anos"] == 2 and r["conclusao"] == "NAO_CONSUMADA"


def test_prescricao_veda_retroativa_antes_da_denuncia():
    e = {"pena": "6m", "modalidade": "concreto", "transito_acusacao": True, "data_fato": "2015-01-01",
         "marcos": [{"evento": "recebimento da denúncia", "data": "2019-01-01"}],
         "data_referencia": "2020-01-01"}
    r = prescricao.calcular(e)
    assert r["intervalos"][0]["consumada"] is False and "12.234" in r["intervalos"][0]["observacao"]


# ---------- minutas ----------
@pytest.fixture
def voto_anotado(tmp_path):
    src = tmp_path / "voto.txt"
    src.write_text(VOTO_OK, encoding="utf-8")
    out = tmp_path / "Minuta_voto.docx"
    montar_minuta.montar(src.read_text(encoding="utf-8").splitlines(), out, montar_minuta.carregar_config(None))
    return out


def test_portao_aprova_voto_bem_formado(voto_anotado):
    r = verificar_minuta.verificar(voto_anotado, "voto", "anotada")
    assert r["bloqueantes"] == []


def test_versao_limpa_sem_vermelho_e_rtf(voto_anotado):
    assert gerar_versoes.main([str(voto_anotado), "--tipo", "voto"]) == 0
    limpa = voto_anotado.with_name("Minuta_voto_LIMPA.docx")
    rtf = limpa.with_suffix(".rtf")
    assert rtf.exists() and rtf.read_text(encoding="ascii").startswith("{\\rtf1")
    texto = "\n".join(p.text for p in __import__("docx").Document(limpa).paragraphs)
    assert "Conferir" not in texto and "615/2025" not in texto and "conferir certidão" not in texto
    assert texto.strip().endswith("É como voto.")


def test_portao_bloqueia_vicios(tmp_path):
    ruim = tmp_path / "ruim.txt"
    ruim.write_text("# DECISÃO MONOCRÁTICA\nTrata-se de habeas corpus impetrado em 12/03/2026 em favor de **A. B.** contra ato do juízo.\n"
                    "Brevemente relatado, passo a decidir.\nI - Da preliminar\nRejeito a preliminar.\n"
                    "A leitura integral dos autos revela a ilegalidade apontada pelo impetrante.\nPublique-se. Intime-se.\n",
                    encoding="utf-8")
    out = tmp_path / "ruim.docx"
    montar_minuta.montar(ruim.read_text(encoding="utf-8").splitlines(), out, montar_minuta.carregar_config(None))
    b = " | ".join(verificar_minuta.verificar(out, "decisao", "anotada")["bloqueantes"])
    for esperado in ("publique-se", "data ou horário no relatório", "numeração de seção",
                     "fragmentad", "linguagem de método", "última linha"):
        assert esperado in b, esperado
    assert "'ocr'" not in b  # "monocrática" não pode disparar falso positivo


# ---------- citações ----------
def test_extracao_de_citacoes():
    t = ("Conforme o AgRg no HC 598.051/SP e a Súmula 231/STJ, aplica-se o Tema 1.139/STJ, "
         "nos termos dos arts. 59 e 68 do CP e do art. 617 do Código de Processo Penal, bem como do "
         "art. 33 da Lei n.º 11.343/2006 e do art. 5º do Regimento Interno do TJAL.")
    chaves = {a["chave_norm"] for a in conferir_citacoes.extrair(t)}
    assert {"AGRG NO HC 598051", "SUMULA STJ 231", "TEMA STJ 1139", "CP ART 59", "CP ART 68",
            "CPP ART 617", "Lei 11343/2006 ART 33", "RITJAL ART 5"} <= chaves


def test_conferencia_contra_ledger(tmp_path, voto_anotado):
    ledger = tmp_path / "_verificacoes.json"
    ledger.write_text(json.dumps([
        {"id": "V1", "tipo": "precedente", "chave": "HC 598.051/SP", "status": "VERIFIED"},
        {"id": "V2", "tipo": "dispositivo", "chave": "art. 33 da Lei n.º 11.343/2006", "status": "VERIFIED"},
    ]), encoding="utf-8")
    assert conferir_citacoes.conferir(voto_anotado, ledger)["aprovada"]
    ledger.write_text(json.dumps([{"id": "V1", "tipo": "precedente", "chave": "HC 598.051/SP", "status": "REJECTED"}]),
                      encoding="utf-8")
    r = conferir_citacoes.conferir(voto_anotado, ledger)
    assert not r["aprovada"] and any("REJEITADA" in b["problema"] for b in r["bloqueantes"])


# ---------- inventário (somente leitura) ----------
def test_inventario_detecta_novo_e_alterado_sem_escrever_na_pasta(tmp_path, voto_anotado):
    pasta = tmp_path / "compartilhada"
    pasta.mkdir()
    alvo = pasta / "voto_relator.docx"
    alvo.write_bytes(voto_anotado.read_bytes())
    saida = tmp_path / "trabalho" / "_inv.json"
    saida.parent.mkdir()
    script = SKILL / "scripts" / "inventario_votos.py"
    antes = sorted(p.name for p in pasta.iterdir())
    subprocess.run([sys.executable, str(script), str(pasta), "--saida", str(saida)], check=True, capture_output=True)
    inv = json.loads(saida.read_text(encoding="utf-8"))
    assert inv["itens"][0]["situacao"] == "NOVO" and inv["itens"][0]["tipo"] in ("relatório e voto", "voto")
    with open(alvo, "ab") as f:
        f.write(b"\0")
    subprocess.run([sys.executable, str(script), str(pasta), "--saida", str(saida)], check=True, capture_output=True)
    assert json.loads(saida.read_text(encoding="utf-8"))["itens"][0]["situacao"] == "ALTERADO"
    assert sorted(p.name for p in pasta.iterdir()) == antes


# ---------- regimento ----------
def test_regimento_carrega_todos_os_artigos():
    artigos = regimento.carregar()
    assert set(range(1, 396)) <= set(artigos)


def test_regimento_texto_vigente_sem_revogados():
    art93 = regimento.consultar([93])[0]
    assert "15 dias" in art93["texto"] and "noventa dias" not in art93["texto"]
    art63 = regimento.consultar([63])[0]["texto"]
    assert "Parágrafo único. Estando impossibilitado" not in art63 and "§ 1º Estando impossibilitado" in art63
    art32 = regimento.consultar([32])[0]["texto"]
    assert "Parágrafo Único. Quando na Câmara Criminal" not in art32 and "§ 1º Quando na Câmara Criminal" in art32
    revogados = (SKILL / "referencias" / "ritjal_revogados.txt").read_text(encoding="utf-8")
    assert "Parágrafo único. Estando impossibilitado" in revogados and "noventa dias" in revogados
    assert all(len(v) == 1 for v in regimento.carregar().values())  # nenhuma redação duplicada


def test_regimento_preserva_hifens_reais():
    art61 = regimento.consultar([61])[0]["texto"]
    assert "assiná-los" in art61 and "Procurador(a)-Geral" in art61
    assert "competindo-lhe" in regimento.consultar([49])[0]["texto"]
    assert regimento.unir(["o(a) Procurador(a)-", "Geral de Justiça", "de 2005 -", "Código"]) == \
        "o(a) Procurador(a)-Geral de Justiça de 2005 - Código"


def test_extracao_do_pdf_reproduz_o_texto_do_skill(tmp_path):
    pytest.importorskip("pdfplumber")
    import extrair_regimento
    pdf = SKILL / "referencias" / "ritjal_consolidado_emenda19.pdf"
    assert extrair_regimento.main([str(pdf), "--saida", str(tmp_path)]) == 0

    def corpo(arquivo):
        return [l for l in arquivo.read_text(encoding="utf-8").splitlines() if not l.startswith("#")]
    assert corpo(tmp_path / "ritjal_integral.txt") == corpo(SKILL / "referencias" / "ritjal_integral.txt")
    assert corpo(tmp_path / "ritjal_revogados.txt") == corpo(SKILL / "referencias" / "ritjal_revogados.txt")


def test_regimento_nao_corta_artigo_em_linha_que_comeca_por_secao():
    # art. 61, XIII, quebra a linha em "Seção Especializada Cível ou nas Câmaras..."
    art61 = regimento.consultar([61])[0]["texto"]
    assert "XXIII - exercer o controle" in art61


def test_regimento_referendo_liminar_camara_criminal():
    art63 = regimento.consultar([63])[0]["texto"]
    assert "decisões concessivas proferidas em feitos da competência da Câmara Criminal" in art63


def test_regimento_busca_e_fila(tmp_path):
    assert any(r["artigo"] == 179 for r in regimento.buscar("setenta e duas horas"))
    src = tmp_path / "m.txt"
    src.write_text("# VOTO\nNos termos do art. 62 do Regimento Interno deste Tribunal e do art. 34 do RISTJ.\n", encoding="utf-8")
    out = tmp_path / "m.docx"
    montar_minuta.montar(src.read_text(encoding="utf-8").splitlines(), out, montar_minuta.carregar_config(None))
    fila = regimento.fila(out)
    assert [e["id"] for e in fila] == ["R062"] and fila[0]["status"] == "PENDENTE"


def test_citacao_regimento_stj_nao_vira_ritjal():
    chaves = {a["chave_norm"] for a in conferir_citacoes.extrair(
        "art. 34 do Regimento Interno do STJ e art. 21 do Regimento Interno do STF, art. 63 deste Regimento")}
    assert {"RISTJ ART 34", "RISTF ART 21", "RITJAL ART 63"} <= chaves
    assert "RITJAL ART 34" not in chaves


def test_portao_aceita_voto_vencido_e_referendo(tmp_path):
    for tipo, titulo in (("voto_vencido", "VOTO VENCIDO"), ("referendo", "VOTO (REFERENDO DE LIMINAR)")):
        src = tmp_path / f"{tipo}.txt"
        src.write_text(f"# {titulo}\nCom a devida vênia ao eminente relator, divirjo quanto à dosimetria, porque a "
                       "fração aplicada na terceira fase carece de fundamentação concreta, como demonstram as fls. 210.\n"
                       "É como voto.\n", encoding="utf-8")
        out = tmp_path / f"{tipo}.docx"
        montar_minuta.montar(src.read_text(encoding="utf-8").splitlines(), out, montar_minuta.carregar_config(None))
        assert verificar_minuta.verificar(out, tipo, "anotada")["bloqueantes"] == [], tipo


# ---------- revisão final: falsos positivos, formatos e modos ----------
def _minuta(tmp_path, nome, texto):
    src = tmp_path / f"{nome}.txt"
    src.write_text(texto, encoding="utf-8")
    out = tmp_path / f"{nome}.docx"
    montar_minuta.montar(src.read_text(encoding="utf-8").splitlines(), out, montar_minuta.carregar_config(None))
    return out


def test_portao_nao_confunde_fatos_criminais_com_linguagem_de_metodo(tmp_path):
    out = _minuta(tmp_path, "vogal", "# VOTO DIVERGENTE\n"
        "Com a devida vênia ao eminente Desembargador Relator, divirjo quanto à validade da prova, porque a "
        "extração de dados do celular apreendido às fls. 40 ocorreu sem autorização judicial, e a varredura "
        "realizada no imóvel não foi precedida de fundadas razões, como se lê das fls. 12/15.\n"
        "Ademais, a distância percorrida pelo veículo, registrada às fls. 22, não corrobora a versão "
        "acusatória, conforme a lição de que \"a prova [...] deve ser lícita\", segundo o laudo de fls. 30.\n"
        "Ante o exposto, voto no sentido de dar provimento ao recurso, acompanhando, no mais, o eminente "
        "Desembargador Relator quanto à rejeição das demais preliminares suscitadas pela defesa.\n"
        "É como voto.\n")
    r = verificar_minuta.verificar(out, "voto_vogal", "anotada")
    assert r["bloqueantes"] == [], r["bloqueantes"]


def test_portao_admite_colchete_de_supressao_na_limpa(voto_anotado, tmp_path):
    out = _minuta(tmp_path, "dec", "# DESPACHO\n"
        "Conforme a orientação segundo a qual \"o réu [...] deve ser intimado pessoalmente\", intime-o "
        "pessoalmente para constituir novo defensor no prazo legal, com a advertência de nomeação de dativo.\n"
        "Publicações e intimações via DJEN\n")
    gerar_versoes.limpar(out, tmp_path / "dec_LIMPA.docx")
    r = verificar_minuta.verificar(tmp_path / "dec_LIMPA.docx", "despacho", "limpa")
    assert not any("colchete" in b for b in r["bloqueantes"]), r["bloqueantes"]


def test_portao_rejeita_tipo_desconhecido(voto_anotado):
    assert verificar_minuta.main([str(voto_anotado), "--tipo", "voto-vista", "--versao", "anotada"]) == 2


def test_rtf_ida_e_volta_preserva_texto_e_citacoes(voto_anotado):
    assert gerar_versoes.main([str(voto_anotado), "--tipo", "voto"]) == 0
    rtf = voto_anotado.with_name("Minuta_voto_LIMPA.rtf")
    texto = conferir_citacoes.texto_de(rtf)
    assert "É como voto." in texto and "apelação criminal" in texto
    assert "HC 598051" in {a["chave_norm"] for a in conferir_citacoes.extrair(texto)}
    especial = "linha 1\nlinha 2\tfim " + chr(0x1F600) + " ç"
    assert conferir_citacoes.rtf_para_texto("{\\rtf1 " + gerar_versoes._esc(especial) + "}") == especial


def test_odt_extracao_nativa(tmp_path):
    import zipfile
    odt = tmp_path / "voto.odt"
    with zipfile.ZipFile(odt, "w") as z:
        z.writestr("content.xml", '<?xml version="1.0"?><office:document-content xmlns:office="o" '
                   'xmlns:text="t"><office:body><text:p>Nos termos do art. 62 do RITJAL,'
                   '<text:s text:c="2"/>nego provimento.</text:p></office:body></office:document-content>')
    texto = conferir_citacoes.texto_de(odt)
    assert "art. 62 do RITJAL,  nego provimento." in texto  # <text:s text:c="2"/> = dois espaços


def test_dosimetria_modo_somado_e_sucessivo():
    base = {"pena_min": "6a", "pena_max": "20a", "fase1": {"desfavoraveis": 0},
            "fase2": {"agravantes": 2, "atenuantes": 0, "fracao": "1/6"}}
    assert dosimetria.calcular(base)["pena_final"] == "8 anos e 2 meses"
    somado = dict(base, fase2=dict(base["fase2"], modo="somado"))
    assert dosimetria.calcular(somado)["pena_final"] == "8 anos"


def test_prescricao_recusa_datas_fora_de_ordem():
    with pytest.raises(ValueError):
        prescricao.calcular({"pena": "2a", "modalidade": "abstrato", "data_fato": "2020-01-10",
                             "marcos": [{"evento": "recebimento da denúncia", "data": "2019-01-01"}],
                             "data_referencia": "2026-10-01"})


# ---------- revisão adversarial dos scripts: regressões ----------
def test_prescricao_art109_vi_fato_anterior_a_lei_12234():
    e = {"pena": "6m", "modalidade": "concreto", "transito_acusacao": True, "data_fato": "2009-03-01",
         "marcos": [{"evento": "recebimento da denúncia", "data": "2009-06-01"},
                    {"evento": "publicação da sentença condenatória", "data": "2011-08-01"}],
         "data_referencia": "2011-09-01"}
    r = prescricao.calcular(e)
    assert r["prazo_anos"] == 2 and r["conclusao"] == "PRESCRICAO_CONSUMADA"
    from datetime import date
    seis_meses = prescricao.pena_exata("6m")
    assert prescricao.prazo_art109(seis_meses, date(2010, 5, 5)) == 2
    assert prescricao.prazo_art109(seis_meses, date(2010, 5, 6)) == 3  # vigência da Lei n.º 12.234/2010


def test_prescricao_termo_no_dia_inexistente_vai_ao_fim_do_mes():
    e = {"pena": "1a4m", "modalidade": "concreto", "transito_acusacao": True, "menor_21_no_fato": True,
         "marcos": [{"evento": "recebimento da denúncia", "data": "2016-02-29"},
                    {"evento": "publicação da sentença condenatória", "data": "2018-02-28"}],
         "data_referencia": "2018-02-28"}
    r = prescricao.calcular(e)
    assert r["intervalos"][0]["termo_final_prazo"] == "2018-02-28" and r["conclusao"] == "NAO_CONSUMADA"


def test_prescricao_suspensoes_sobrepostas_e_posteriores_ao_termo():
    base = {"pena": "4a", "modalidade": "abstrato", "data_referencia": "2019-03-01",
            "marcos": [{"evento": "recebimento da denúncia", "data": "2010-01-01"}]}
    sobrepostas = prescricao.calcular({**base, "suspensoes": [{"inicio": "2012-01-01", "fim": "2012-12-31"},
                                                              {"inicio": "2012-06-01", "fim": "2012-12-31"}]})
    assert sobrepostas["intervalos"][0]["dias_suspensos"] == 366
    assert sobrepostas["conclusao"] == "PRESCRICAO_CONSUMADA"
    # suspensão iniciada no dia seguinte ao esgotamento do prazo não o reabre
    tardia = prescricao.calcular({**base, "data_referencia": "2018-12-31",
                                  "suspensoes": [{"inicio": "2018-01-01", "fim": "2018-12-31"}]})
    assert tardia["conclusao"] == "PRESCRICAO_CONSUMADA"


def test_prescricao_modalidade_ausente_equivale_a_concreto():
    e = {"pena": "6m", "transito_acusacao": True, "data_fato": "2015-01-01",
         "marcos": [{"evento": "recebimento da denúncia", "data": "2019-01-01"}], "data_referencia": "2020-01-01"}
    sem = prescricao.calcular(e)
    com = prescricao.calcular({**e, "modalidade": "concreto"})
    assert sem["conclusao"] == com["conclusao"] == "NAO_CONSUMADA"
    assert any("não informada" in o for o in sem["observacoes"])
    with pytest.raises(ValueError):
        prescricao.calcular({**e, "modalidade": "retroativa"})


def test_dosimetria_valida_tipo_e_fracao_da_terceira_fase():
    base = {"pena_min": "5a", "pena_max": "15a"}
    assert dosimetria.calcular({**base, "fase3": [{"tipo": "Aumento", "fracao": "1/3"}]})["pena_final"] == \
        "6 anos e 8 meses"
    assert dosimetria.calcular({**base, "fase3": [{"tipo": "diminuição", "fracao": "1/3"}]})["pena_final"] == \
        "3 anos e 4 meses"
    for ruim in ({"tipo": "majorante", "fracao": "1/3"}, {"tipo": "diminuicao", "fracao": "3/2"}):
        with pytest.raises(ValueError):
            dosimetria.calcular({**base, "fase3": [ruim]})
    with pytest.raises(ValueError):
        dosimetria.calcular({**base, "multa_min": 10})


def test_dosimetria_despreza_fracao_de_dia_em_cada_fase():
    e = {"pena_min": "1a", "pena_max": "4a", "fase1": {"desfavoraveis": 1},
         "fase2": {"agravantes": 0, "atenuantes": 1, "fracao": "1/6"},
         "fase3": [{"tipo": "aumento", "fracao": "1/3"}], "afirmado": {"pena_final": "1a6m9d"}}
    r = dosimetria.calcular(e)
    assert r["pena_final"] == "1 ano, 6 meses e 9 dias" and r["conclusao"] == "CONFERE"


def test_dosimetria_art44_crime_culposo():
    r = dosimetria.calcular({"pena_min": "5a", "pena_max": "5a", "culposo": True})
    assert "culposo" in r["quadro_art44"] and "incabível" not in r["quadro_art44"]


def test_citacoes_formas_plurais_grau_cnj_e_inversa():
    t = ("Nos termos das Súmulas 718 e 719 do STF, do HC n.° 598.051/SP, do RESP 1.234.567/SP e dos Temas "
         "1.139 e 1.140/STJ (CF, art. 93, IX), esta Câmara Criminal, na Apelação Criminal n.º "
         "0700123-83.2024.8.02.0001, Rel. Des. Fulano, decidiu no mesmo sentido.")
    chaves = {a["chave_norm"]: a["tipo"] for a in conferir_citacoes.extrair(t)}
    for k in ("SUMULA STF 718", "SUMULA STF 719", "HC 598051", "RESP 1234567", "TEMA STJ 1139",
              "TEMA STJ 1140", "CF ART 93", "CNJ 0700123-83.2024.8.02.0001"):
        assert k in chaves, k
    assert chaves["CNJ 0700123-83.2024.8.02.0001"] == "precedente"
    origem = conferir_citacoes.extrair("nos autos da ação penal n.º 0700456-12.2023.8.02.0058, a defesa requereu")
    assert [a["tipo"] for a in origem] == ["processo"]
    assert conferir_citacoes.extrair(t, processo="0700123-83.2024.8.02.0001") and \
        "CNJ 0700123-83.2024.8.02.0001" not in {a["chave_norm"] for a in conferir_citacoes.extrair(
            t, processo="0700123-83.2024.8.02.0001")}


def test_citacoes_artigo_sem_diploma_nao_e_atribuido_ao_seguinte():
    chaves = {a["chave_norm"] for a in conferir_citacoes.extrair("art. 5º, LVII, e o art. 386, VII, do CPP")}
    assert chaves == {"CPP ART 386", "? ART 5"}


def test_ledger_usa_chave_norm_e_o_status_mais_restritivo(tmp_path):
    minuta = tmp_path / "m.txt"
    minuta.write_text("A pena observa o art. 33 da Lei n.º 11.343/2006 e o art. 59 do CP.", encoding="utf-8")
    fila = [{"id": "P001", "chave": "art. 33 da Lei n.º 11.343/2006 e o art. 59 do CP",
             "chave_norm": "Lei 11343/2006 ART 33", "status": "VERIFIED"},
            {"id": "P002", "chave": "art. 59 do CP", "chave_norm": "CP ART 59", "status": "PENDENTE"}]
    ledger = tmp_path / "l.json"
    ledger.write_text(json.dumps(fila), encoding="utf-8")
    r = conferir_citacoes.conferir(minuta, ledger)
    assert not r["aprovada"] and [b["chave_norm"] for b in r["bloqueantes"]] == ["CP ART 59"]
    m3 = tmp_path / "m3.txt"
    m3.write_text("Aplica-se a Súmula 231/STJ.", encoding="utf-8")
    ledger.write_text(json.dumps([{"id": "V010", "chave": "Súmula 231/STJ", "status": "REJECTED"},
                                  {"id": "V001", "chave": "Súmula 231/STJ", "status": "VERIFIED"}]), encoding="utf-8")
    assert not conferir_citacoes.conferir(m3, ledger)["aprovada"]


def _docx_com_hiperlink_e_tabela(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    out = _minuta(tmp_path, "hl", "# VOTO\nCom a devida vênia, divirjo, porque a jurisprudência é firme, como se vê "
                  "do PRECEDENTE, no sentido de que a minorante deve ser aplicada ao caso concreto dos autos.\n"
                  "Ante o exposto, voto no sentido de dar provimento ao recurso, nos termos da fundamentação.\n"
                  "É como voto.\n{{ressalva}}\n{{advertencia}}\n")
    doc = __import__("docx").Document(out)
    p = [p for p in doc.paragraphs if "PRECEDENTE" in p.text][0]
    r = [r for r in p.runs if "PRECEDENTE" in r.text][0]
    antes, depois = r.text.split("PRECEDENTE")
    r.text = antes
    h = OxmlElement("w:hyperlink")
    hr = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "AgRg no HC 598.051/SP"
    hr.append(t)
    h.append(hr)
    r._r.addnext(h)
    resto = OxmlElement("w:r")
    t2 = OxmlElement("w:t")
    t2.text = depois
    t2.set(qn("xml:space"), "preserve")
    resto.append(t2)
    h.addnext(resto)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Publique-se. Súmula 444/STJ."
    doc.save(out)
    return out


def test_estruturas_nao_suportadas_bloqueiam_e_citacoes_sao_lidas(tmp_path):
    out = _docx_com_hiperlink_e_tabela(tmp_path)
    b = " | ".join(verificar_minuta.verificar(out, "voto", "anotada")["bloqueantes"])
    assert "tabela" in b and "hiperlink" in b
    chaves = {a["chave_norm"] for a in conferir_citacoes.extrair(conferir_citacoes.texto_de(out))}
    assert {"AGRG NO HC 598051", "SUMULA STJ 444"} <= chaves
    assert gerar_versoes.main([str(out), "--tipo", "voto"]) == 1  # nada vai ao RTF


def test_portao_regressoes_de_falso_negativo(tmp_path):
    casos = {
        ("ementa", "linguagem de método"): "# EMENTA\nDIREITO PENAL. APELAÇÃO CRIMINAL. Após leitura integral dos "
                                           "autos e extração do texto por OCR, verificou-se a ausência de prova.\n",
        ("relatorio", "data ou horário"): "# RELATÓRIO\nTrata-se de apelação criminal interposta por J. S. contra a "
                                          "sentença de fls. 210/218, publicada em 1.º de outubro de 2025.\nA audiência "
                                          "realizou-se às 14h, conforme termo de fls. 120/125, com a oitiva de três "
                                          "testemunhas arroladas pela acusação.\nÉ o relatório.\n",
        ("voto_vogal", "período fragmentado"): "# VOTO DIVERGENTE\nCom a devida vênia ao eminente relator, divirjo "
                                               "quanto à dosimetria, porque a fração aplicada carece de fundamentação "
                                               "concreta, como demonstram as fls. 210.\nNego-lhe provimento.\n"
                                               "É como voto.\n",
        ("voto_vogal", "negrito aplicado"): "# VOTO DIVERGENTE\nCom a devida vênia ao eminente relator, divirjo, como "
                                            "demonstram as fls. 210. **O réu deve ser absolvido por insuficiência de "
                                            "provas, nos termos do art. 386, VII, do CPP**.\nÉ como voto.\n",
        ("despacho", "abreviada"): "# DESPACHO\nConsiderando a renúncia do único defensor constituído, intime-se o "
                                   "réu pessoalmente para constituir novo advogado no prazo legal. P. R. I.\n"
                                   "Publicações e intimações via DJEN\n",
    }
    for (tipo, esperado), texto in casos.items():
        b = " | ".join(verificar_minuta.verificar(_minuta(tmp_path, tipo, texto), tipo, "anotada")["bloqueantes"])
        assert esperado in b, (tipo, esperado, b)
    sem_titulo = _minuta(tmp_path, "st", "Trata-se de habeas corpus impetrado em 12/03/2026 em favor de A. B.\n"
                         "Brevemente relatado, passo a decidir.\nPublicações e intimações via DJEN\n")
    assert any("data ou horário" in b for b in verificar_minuta.verificar(sem_titulo, "decisao", "anotada")["bloqueantes"])


def test_portao_regressoes_de_falso_positivo(tmp_path):
    import unicodedata
    texto = ("# VOTO DIVERGENTE\nCom a devida vênia, divirjo, porque o texto extraído do aparelho celular e a mídia "
             "examinada visualmente pelo perito, às fls. 40, não comprovam a autoria [grifo meu], nem o laudo [sic!].\n"
             "> II – o réu não for reincidente em crime doloso;\n"
             "Ante o exposto, acompanho o eminente Desembargador Relator e nego provimento ao recurso.\n"
             "É como voto.\n")
    b = verificar_minuta.verificar(_minuta(tmp_path, "fp", unicodedata.normalize("NFD", texto)),
                                   "voto_vogal", "anotada")["bloqueantes"]
    assert not [x for x in b if "ressalva" not in x and "advertência" not in x and "vermelha" not in x], b
    limpa = tmp_path / "fp_LIMPA.docx"
    gerar_versoes.limpar(_minuta(tmp_path, "fp2", texto), limpa)
    assert not any("colchete" in x for x in verificar_minuta.verificar(limpa, "voto_vogal", "limpa")["bloqueantes"])


def test_versao_limpa_sem_espaco_residual(voto_anotado):
    assert gerar_versoes.main([str(voto_anotado), "--tipo", "voto"]) == 0
    texto = "\n".join(p.text for p in __import__("docx").Document(
        voto_anotado.with_name("Minuta_voto_LIMPA.docx")).paragraphs)
    assert "privilegiado, à luz" in texto and " ," not in texto and "  " not in texto


def test_inventario_protege_a_pasta_e_preserva_o_estado(tmp_path):
    pasta = tmp_path / "compartilhada"
    pasta.mkdir()
    (pasta / "voto.txt").write_text("VOTO\nRelator: Des. Fulano de Tal\n0700123-83.2024.8.02.0001\n", encoding="utf-8")
    script = [sys.executable, str(SKILL / "scripts" / "inventario_votos.py")]
    assert subprocess.run(script + [str(pasta)], capture_output=True, cwd=pasta).returncode == 2
    assert subprocess.run(script + [str(pasta), "--saida", str(pasta / "x.json")], capture_output=True).returncode == 2
    assert sorted(p.name for p in pasta.iterdir()) == ["voto.txt"]
    saida = tmp_path / "inv.json"

    def rodada(alvo=pasta):
        subprocess.run(script + [str(alvo), "--saida", str(saida)], check=True, capture_output=True)
        return json.loads(saida.read_text(encoding="utf-8"))["itens"][0]
    assert rodada()["situacao"] == "NOVO"
    assert rodada(pasta.resolve())["situacao"] == "NOVO"  # caminho absoluto não apaga o histórico
    inv = json.loads(saida.read_text(encoding="utf-8"))
    inv["itens"][0].update(revisado_em="2026-10-01", nota="Revisao_0700123.docx")
    saida.write_text(json.dumps(inv), encoding="utf-8")
    item = rodada()
    assert item["situacao"] == "INALTERADO" and item["nota"] == "Revisao_0700123.docx"
    with open(pasta / "voto.txt", "a", encoding="utf-8") as f:
        f.write("acréscimo\n")
    item = rodada()
    assert item["situacao"] == "ALTERADO" and item["revisao_anterior"] == "2026-10-01" and not item["revisado_em"]
    assert rodada()["situacao"] == "ALTERADO"  # persiste até nova revisão


def test_regimento_e_cnj_sinalizam_entrada_invalida():
    assert regimento.main(["63-A"]) == 1
    assert cnj.main(["--completar", "0700123-00.2024", "--foro", "0001"]) == 1
    assert cnj.main(["--completar", "0700123-83.2024", "--foro", "0001"]) == 0
