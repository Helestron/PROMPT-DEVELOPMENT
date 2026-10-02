#!/usr/bin/env python3
"""Monta a VERSÃO ANOTADA (.docx) de uma minuta a partir de texto com marcação leve.

Marcação aceita (uma linha = um parágrafo; linha em branco é ignorada):
    # TÍTULO                 epígrafe estrutural — centralizada, negrito, caixa alta; só as
                             admitidas pelo portão (verificar_minuta.py): DESPACHO, DECISÃO,
                             DECISÃO MONOCRÁTICA, RELATÓRIO, VOTO, VOTO-VISTA, VOTO DIVERGENTE,
                             VOTO VENCIDO, DECLARAÇÃO DE VOTO, VOTO (REFERENDO DE LIMINAR),
                             EMENTA e NOTA DE REVISÃO
    > texto                  citação longa (recuo de 4 cm, espaçamento simples, fonte menor)
    **trecho**               negrito
    _trecho_                 itálico (só para estrangeirismos)
    {{v: trecho}}            destaque de conferência em VERMELHO e negrito (só na anotada)
    {{conferir: motivo}}     gera "[Conferir: motivo]" em vermelho
    {{ressalva}}             insere a ressalva de apoio à decisão (art. 93, IX, da CF;
                             Resolução CNJ n.º 615/2025)
    {{advertencia}}          insere a advertência final sobre os trechos em vermelho

Uso:
    python montar_minuta.py entrada.txt saida_anotada.docx [--config ../config/gabinete.json]

A ressalva e a advertência são acrescentadas automaticamente ao fim, se o texto não as trouxer.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

VERMELHO = RGBColor(0xFF, 0x00, 0x00)
RESSALVA = ("Minuta de apoio elaborada com auxílio de inteligência artificial, sujeita à revisão, "
            "à alteração e à decisão do magistrado (art. 93, IX, da CF; Resolução CNJ n.º 615/2025). "
            "Uso interno do gabinete.")
ADVERTENCIA = ("Os trechos em vermelho são apontamentos de conferência e devem ser revistos e "
               "removidos antes da assinatura.")
TOKEN = re.compile(r"(\*\*.+?\*\*|_[^_]+?_|\{\{v:\s*.+?\}\}|\{\{conferir:\s*.+?\}\})")


def carregar_config(caminho):
    padrao = {"fonte": "Times New Roman", "corpo_pt": 12, "citacao_pt": 11,
              "recuo_primeira_linha_cm": 2.0, "entrelinhas": 1.5}
    if caminho and Path(caminho).exists():
        padrao.update(json.loads(Path(caminho).read_text(encoding="utf-8")).get("formatacao", {}))
    return padrao


def runs(par, texto, cfg, tamanho):
    for pedaco in TOKEN.split(texto):
        if not pedaco:
            continue
        r = None
        if pedaco.startswith("**") and pedaco.endswith("**"):
            r = par.add_run(pedaco[2:-2]); r.bold = True
        elif pedaco.startswith("_") and pedaco.endswith("_") and len(pedaco) > 2:
            r = par.add_run(pedaco[1:-1]); r.italic = True
        elif pedaco.startswith("{{v:"):
            r = par.add_run(pedaco[4:-2].strip()); r.bold = True; r.font.color.rgb = VERMELHO
        elif pedaco.startswith("{{conferir:"):
            r = par.add_run(f"[Conferir: {pedaco[11:-2].strip()}]"); r.bold = True; r.font.color.rgb = VERMELHO
        else:
            r = par.add_run(pedaco)
        r.font.name = cfg["fonte"]; r.font.size = Pt(tamanho)


def paragrafo_vermelho(doc, texto, cfg):
    p = doc.add_paragraph()
    r = p.add_run(texto); r.bold = True; r.font.color.rgb = VERMELHO
    r.font.name = cfg["fonte"]; r.font.size = Pt(cfg["corpo_pt"])
    p.paragraph_format.space_before = Pt(12)
    return p


def montar(linhas, saida, cfg):
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = cfg["fonte"]; estilo.font.size = Pt(cfg["corpo_pt"])
    tem_ressalva = tem_adv = False
    for bruta in linhas:
        bruta = unicodedata.normalize("NFC", bruta)
        linha = bruta.rstrip("\n")
        if not linha.strip():
            continue
        if linha.strip() == "{{ressalva}}":
            paragrafo_vermelho(doc, RESSALVA, cfg); tem_ressalva = True; continue
        if linha.strip() == "{{advertencia}}":
            paragrafo_vermelho(doc, ADVERTENCIA, cfg); tem_adv = True; continue
        if linha.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(linha[2:].strip().upper()); r.bold = True
            r.font.name = cfg["fonte"]; r.font.size = Pt(cfg["corpo_pt"])
            p.paragraph_format.space_after = Pt(12)
            continue
        if linha.startswith("> "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(4); p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(12)
            runs(p, linha[2:], cfg, cfg["citacao_pt"])
            continue
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(cfg["recuo_primeira_linha_cm"])
        p.paragraph_format.line_spacing = cfg["entrelinhas"]
        p.paragraph_format.space_after = Pt(6)
        runs(p, linha, cfg, cfg["corpo_pt"])
    if not tem_ressalva:
        paragrafo_vermelho(doc, RESSALVA, cfg)
    if not tem_adv:
        paragrafo_vermelho(doc, ADVERTENCIA, cfg)
    doc.save(saida)


def main(argv):
    if len(argv) < 2:
        print(__doc__); return 2
    cfg_path = argv[argv.index("--config") + 1] if "--config" in argv else \
        Path(__file__).resolve().parent.parent / "config" / "gabinete.json"
    cfg = carregar_config(cfg_path)
    montar(Path(argv[0]).read_text(encoding="utf-8").splitlines(), argv[1], cfg)
    print(f"Versão anotada gravada em {argv[1]}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
