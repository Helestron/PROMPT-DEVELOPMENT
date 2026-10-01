#!/usr/bin/env python3
"""Gera a VERSÃO LIMPA (.docx e .rtf) a partir da versão anotada, sob o portão de verificação.

Passos:
  1. remove todo run em vermelho (FF0000) — apontamentos de conferência, ressalva e advertência;
  2. remove parágrafos que fiquem vazios após a remoção;
  3. grava <base>_LIMPA.docx e roda verificar_minuta.py --versao limpa sobre ele;
  4. só havendo aprovação (sem pendência bloqueante), converte para .rtf — motor nativo
     (padrão, sem dependências), LibreOffice (--motor soffice) ou, no Windows, Word COM
     (saj_sg5.ps1 -Modo ConverterRtf).

Uso:
    python gerar_versoes.py Minuta_<numero>_<ato>.docx --tipo voto|voto_vogal|decisao|despacho|ementa
                            [--motor nativo|soffice]
Saída: caminho do .rtf aprovado, ou relatório de pendências (código de saída 1).
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
import verificar_minuta  # noqa: E402


def eh_vermelho(run) -> bool:
    cor = run.font.color
    return cor is not None and cor.type is not None and cor.rgb is not None and str(cor.rgb).upper() == "FF0000"


def limpar(origem: Path, destino: Path):
    doc = Document(origem)
    for p in list(doc.paragraphs):
        for r in list(p.runs):
            if eh_vermelho(r):
                r._element.getparent().remove(r._element)
        if not p.text.strip():
            p._element.getparent().remove(p._element)
    doc.save(destino)


def _esc(texto: str) -> str:
    out = []
    for ch in texto:
        o = ord(ch)
        if ch in "\\{}":
            out.append("\\" + ch)
        elif o < 128:
            out.append(ch)
        else:
            out.append(f"\\u{o if o < 32768 else o - 65536}?")
    return "".join(out)


def _twips(medida) -> int:
    return int(medida.twips) if medida is not None else 0


def docx_para_rtf(docx: Path, destino: Path) -> Path:
    """Conversor nativo DOCX -> RTF (parágrafos, alinhamento, recuos, entrelinhas, negrito,
    itálico, sublinhado, tamanho e cor). Basta ao emissor do SAJ, que recebe texto formatado."""
    doc = Document(docx)
    fonte = doc.styles["Normal"].font.name or "Times New Roman"
    tam_padrao = doc.styles["Normal"].font.size.pt if doc.styles["Normal"].font.size else 12
    partes = ["{\\rtf1\\ansi\\ansicpg1252\\deff0\\uc1",
              "{\\fonttbl{\\f0\\froman " + _esc(fonte) + ";}}",
              "{\\colortbl;\\red0\\green0\\blue0;\\red255\\green0\\blue0;}",
              "\\paperw11906\\paperh16838\\margl1701\\margr1134\\margt1701\\margb1134\n"]
    alinh = {0: "\\ql", 1: "\\qc", 2: "\\qr", 3: "\\qj"}
    for p in doc.paragraphs:
        pf = p.paragraph_format
        cab = "\\pard\\plain" + alinh.get(int(p.alignment) if p.alignment is not None else 0, "\\ql")
        cab += f"\\fi{_twips(pf.first_line_indent)}\\li{_twips(pf.left_indent)}"
        cab += f"\\sa{_twips(pf.space_after)}\\sb{_twips(pf.space_before)}"
        if isinstance(pf.line_spacing, float):
            cab += f"\\sl{int(240 * pf.line_spacing)}\\slmult1"
        corpo = []
        for r in p.runs:
            if not r.text:
                continue
            fmt = "\\f0\\fs" + str(int(2 * (r.font.size.pt if r.font.size else tam_padrao)))
            if r.bold:
                fmt += "\\b"
            if r.italic:
                fmt += "\\i"
            if r.underline:
                fmt += "\\ul"
            cor = r.font.color
            if cor is not None and cor.type is not None and cor.rgb is not None and str(cor.rgb).upper() == "FF0000":
                fmt += "\\cf2"
            corpo.append("{" + fmt + " " + _esc(r.text) + "}")
        partes.append(cab + " " + "".join(corpo) + "\\par\n")
    partes.append("}")
    destino.write_text("".join(partes), encoding="ascii")
    return destino


def para_rtf(docx: Path, motor: str = "nativo") -> Path:
    rtf = docx.with_suffix(".rtf")
    if motor == "nativo":
        return docx_para_rtf(docx, rtf)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice ausente: use o motor nativo ou saj_sg5.ps1 -Modo ConverterRtf (Word COM).")
    subprocess.run([soffice, "--headless", "--convert-to", "rtf", "--outdir", str(docx.parent), str(docx)],
                   check=True, capture_output=True, timeout=180)
    if not rtf.exists():
        raise RuntimeError("Conversão para RTF não produziu arquivo.")
    return rtf


def main(argv):
    if not argv:
        print(__doc__); return 2
    origem = Path(argv[0])
    tipo = argv[argv.index("--tipo") + 1] if "--tipo" in argv else "decisao"
    # A versão anotada também passa pelo portão (presença de ressalva, advertência e cor).
    rel_anot = verificar_minuta.verificar(origem, tipo, "anotada")
    limpa = origem.with_name(origem.stem + "_LIMPA.docx")
    limpar(origem, limpa)
    rel = verificar_minuta.verificar(limpa, tipo, "limpa")
    bloqueios = rel_anot["bloqueantes"] + rel["bloqueantes"]
    if bloqueios:
        print(json.dumps({"aprovada": False, "bloqueantes": bloqueios,
                          "apontamentos": rel_anot["apontamentos"] + rel["apontamentos"]},
                         ensure_ascii=False, indent=2))
        return 1
    motor = argv[argv.index("--motor") + 1] if "--motor" in argv else "nativo"
    rtf = para_rtf(limpa, motor)
    print(json.dumps({"aprovada": True, "rtf": str(rtf), "docx_limpo": str(limpa),
                      "apontamentos": rel["apontamentos"]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
